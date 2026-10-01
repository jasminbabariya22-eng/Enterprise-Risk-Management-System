from sqlalchemy.orm import Session, joinedload
from datetime import datetime, timezone
from sqlalchemy.inspection import inspect
from sqlalchemy import func
import pandas as pd
import io
from sqlalchemy import or_, and_
from datetime import datetime

import math
from openpyxl.styles import PatternFill, Alignment,Font

from app.models.department import Department
from app.models.risk_register import RiskRegister
from app.models.risk_description import RiskDescription
from app.models.risk_treatment import RiskTreatment
from app.models.risk_action_followup import RiskActionFollowup

from app.models.risk_register_hist import RiskRegisterHist
from app.models.risk_description_hist import RiskDescriptionHist
from app.models.risk_treatment_hist import RiskTreatmentHist

from app.models.mst_status import Status
from app.models.user import User

from app.services.email_event_service import send_risk_created_email

# Generate Risk ID
def generate_risk_id(db: Session, dept_id: int):
    dept = db.query(Department).filter(
        Department.id == dept_id
    ).with_for_update().first()

    if not dept:
        raise Exception("Department not found")

    if dept.last_risk_number is None:
        dept.last_risk_number = 0
    dept.last_risk_number += 1
    number = dept.last_risk_number

    short_name = dept.dept_short_name or "RISK"
    risk_id = f"{short_name}-{str(number).zfill(4)}"

    return risk_id


# Type Conversion
def to_int(val):
    try:
        return int(val) if val not in [None, ""] else None
    except (ValueError, TypeError):
        return None


def to_float(val):
    try:
        return float(val) if val not in [None, ""] else None
    except (ValueError, TypeError):
        return None


def to_datetime(val):
    if val in [None, ""]:
        return None
    try:
        if isinstance(val, datetime):
            return val
        return datetime.fromisoformat(str(val))
    except Exception:
        return None


def model_to_dict(obj):
    return {c.key: getattr(obj, c.key) for c in inspect(obj).mapper.column_attrs}


# Get Status ID
def get_status_id(db: Session, status_name: str):
    try:
        status = db.query(Status).filter(
            func.lower(Status.status_name) == func.lower(status_name),
            Status.is_deleted == 0
        ).first()

        if not status:
            status = db.query(Status).filter(
                Status.status_name.ilike(f"%{status_name}%"),
                Status.is_deleted == 0
            ).first()

        return status.id if status else None
    except Exception:
        return None


def reset_risk_approvals(risk):
    risk.risk_function_head_approval_status = None
    risk.risk_function_head_approval_remark = None
    risk.risk_function_head_approval_by = None
    risk.risk_function_head_approval_on = None

    risk.risk_manager_approval_status = None
    risk.risk_manager_approval_remark = None
    risk.risk_manager_approval_by = None
    risk.risk_manager_approved_on = None

    risk.risk_head_approval_status = None
    risk.risk_head_approval_remark = None
    risk.risk_head_approval_by = None
    risk.risk_head_approved_on = None

    return risk

# ---------
# CREATE OR UPDATE RISK
# ---------

def create_update_risk(db: Session, data, current_user):

    try:

        if not data or not data.risk_register:
            raise ValueError("risk_register is required")

        register_data = data.risk_register
        desc_data = data.risk_description
        treatments = data.risk_treatments or []
        user_type_name = str(current_user.get('user_type_name') or current_user.get('user_type') or '')
        
        now_dt = datetime.now(timezone.utc)
        pending_for_action = get_status_id(db, "Pending for Action")
        opened_status = get_status_id(db, "Open") or get_status_id(db, "Opened") or get_status_id(db, "New")

        co_owner = to_int(register_data.risk_co_owner_id)
        if co_owner is not None and co_owner <= 0:
            co_owner = None

        r_status = to_int(register_data.risk_status)
        if not r_status or r_status <= 0:
            draft_id = get_status_id(db, "Draft") or 1
            submitted_id = get_status_id(db, "Submitted to Function Head") or 2
            r_status = submitted_id if user_type_name.upper() == 'RISK OWNER' else draft_id

        # CREATE OR UPDATE RISK REGISTER

        if to_int(register_data.risk_register_id) == 0:

            risk_id = generate_risk_id(db, to_int(register_data.dept_id))
            max_reg_id = db.query(func.max(RiskRegister.risk_register_id)).scalar() or 0
            new_reg_id = max_reg_id + 1

            risk = RiskRegister(
                risk_register_id=new_reg_id,
                risk_id=risk_id,
                risk_name=register_data.risk_name,
                dept_id=to_int(register_data.dept_id),
                risk_owner_id=to_int(register_data.risk_owner_id),
                risk_co_owner_id=co_owner,
                financial_year=register_data.financial_year,
                risk_status=r_status,
                risk_progress=register_data.risk_progress or "0",
                created_by=current_user.get("id", 1),
                created_on=now_dt,
                modified_by=current_user.get("id", 1),
                modified_on=now_dt,
                is_active=0,
                is_deleted=0
            )

            db.add(risk)
            db.flush()
            if opened_status and risk.risk_status == opened_status and user_type_name.upper() == 'RISK OWNER':
                try:
                    send_risk_created_email(db, risk.risk_register_id)
                except Exception as ex_mail:
                    pass

        else:

            risk = db.query(RiskRegister).filter(
                RiskRegister.risk_register_id == to_int(register_data.risk_register_id)
            ).first()

            if not risk:
                raise ValueError("RiskRegister not found")
            
            if user_type_name.upper() == 'RISK OWNER':
                risk.risk_head_approval_by = None
                risk.risk_head_approval_remark = None
                risk.risk_head_approval_status = None
                risk.risk_head_approved_on = None

                risk.risk_manager_approval_by = None
                risk.risk_manager_approval_remark = None
                risk.risk_manager_approval_status = None
                risk.risk_manager_approved_on = None

                risk.function_head_status = None
                risk.risk_function_head_approval_by = None
                risk.risk_function_head_approval_on = None
                risk.risk_function_head_approval_remark = None

            risk.risk_name = register_data.risk_name
            risk.dept_id = to_int(register_data.dept_id)
            risk.risk_owner_id = to_int(register_data.risk_owner_id)
            risk.risk_co_owner_id = co_owner
            risk.financial_year = register_data.financial_year
            if to_int(register_data.risk_status):
                risk.risk_status = to_int(register_data.risk_status)
            risk.risk_progress = register_data.risk_progress or "0"

            risk.modified_by = current_user.get("id", 1)
            risk.modified_on = now_dt
            
            if opened_status and risk.risk_status == opened_status and user_type_name.upper() == "RISK OWNER":
                try:
                    send_risk_created_email(db, risk.risk_register_id)
                except Exception as ex_mail:
                    pass
            
        # HISTORY - RISK REGISTER

        hist_register = RiskRegisterHist(
            risk_register_id=risk.risk_register_id,
            risk_id=risk.risk_id,
            risk_name=risk.risk_name,
            dept_id=risk.dept_id,
            risk_owner_id=risk.risk_owner_id,
            risk_co_owner_id=co_owner,
            financial_year=risk.financial_year,
            risk_status=risk.risk_status,
            risk_progress=risk.risk_progress,
            created_by=risk.created_by or current_user.get("id", 1),
            created_on=risk.created_on or now_dt,
            modified_by=risk.modified_by or current_user.get("id", 1),
            modified_on=risk.modified_on or now_dt,
            is_active=risk.is_active or 0,
            is_deleted=risk.is_deleted or 0
        )

        db.add(hist_register)


        # RISK DESCRIPTION

        description = None

        if desc_data and any([
            desc_data.risk_description not in [None, ""],
            desc_data.mitigation not in [None, ""],
            to_int(desc_data.inherent_risk_likelihood_id) not in [None, 0],
            to_int(desc_data.inherent_risk_impact_id) not in [None, 0],
            to_int(desc_data.current_risk_likelihood_id) not in [None, 0],
            to_int(desc_data.current_risk_impact_id) not in [None, 0]
        ]):

            if to_int(desc_data.risk_description_id) == 0:
                max_desc_id = db.query(func.max(RiskDescription.risk_description_id)).scalar() or 0
                new_desc_id = max_desc_id + 1

                description = RiskDescription(
                    risk_description_id=new_desc_id,
                    risk_register_id=risk.risk_register_id,
                    risk_id=risk.risk_id,
                    risk_description=desc_data.risk_description,
                    inherent_risk_likelihood_id=to_int(desc_data.inherent_risk_likelihood_id),
                    inherent_risk_impact_id=to_int(desc_data.inherent_risk_impact_id),
                    mitigation=desc_data.mitigation,
                    current_risk_likelihood_id=to_int(desc_data.current_risk_likelihood_id),
                    current_risk_impact_id=to_int(desc_data.current_risk_impact_id),
                    created_by=current_user.get("id", 1),
                    created_on=now_dt,
                    modified_by=current_user.get("id", 1),
                    modified_on=now_dt,
                    is_deleted=0
                )

                db.add(description)
                db.flush()

            # Update Risk Description
            else:

                description = db.query(RiskDescription).filter(
                    RiskDescription.risk_description_id == to_int(desc_data.risk_description_id)
                ).first()

                if not description:
                    raise ValueError("RiskDescription not found")

                description.risk_description = desc_data.risk_description
                description.inherent_risk_likelihood_id = to_int(desc_data.inherent_risk_likelihood_id)
                description.inherent_risk_impact_id = to_int(desc_data.inherent_risk_impact_id)
                description.mitigation = desc_data.mitigation
                description.current_risk_likelihood_id = to_int(desc_data.current_risk_likelihood_id)
                description.current_risk_impact_id = to_int(desc_data.current_risk_impact_id)

                description.modified_by = current_user.get("id", 1)
                description.modified_on = now_dt


            # HISTORY DESCRIPTION

            hist_desc = RiskDescriptionHist(
                risk_description_id=description.risk_description_id,
                risk_register_id=description.risk_register_id,
                risk_id=description.risk_id,
                risk_description=description.risk_description,
                inherent_risk_likelihood_id=description.inherent_risk_likelihood_id,
                inherent_risk_impact_id=description.inherent_risk_impact_id,
                mitigation=description.mitigation,
                current_risk_likelihood_id=description.current_risk_likelihood_id,
                current_risk_impact_id=description.current_risk_impact_id,
                created_by=description.created_by or current_user.get("id", 1),
                created_on=description.created_on or now_dt,
                modified_by=description.modified_by or current_user.get("id", 1),
                modified_on=description.modified_on or now_dt,
                is_deleted=description.is_deleted or 0
            )

            db.add(hist_desc)

        
        # RISK TREATMENTS

        saved_treatments = []

        if description is not None:

            if desc_data and to_int(desc_data.risk_description_id) > 0:
                
                db.query(RiskTreatment).filter(
                    RiskTreatment.risk_description_id == description.risk_description_id
                ).delete()

            max_treat_id = db.query(func.max(RiskTreatment.risk_treatment_id)).scalar() or 0

            for t_idx, treatment in enumerate(treatments):
                act_owner = to_int(treatment.action_owner_id)
                if act_owner is None or act_owner <= 0:
                    act_owner = to_int(register_data.risk_owner_id) or current_user.get("id", 1)

                act_status = to_int(treatment.action_status_id)
                if act_status is not None and act_status <= 0:
                    act_status = None

                new_treatment = RiskTreatment(
                    risk_treatment_id=max_treat_id + t_idx + 1,
                    risk_register_id=risk.risk_register_id,
                    risk_description_id=description.risk_description_id,
                    risk_id=risk.risk_id,
                    action_plan=treatment.action_plan,
                    action_owner_id=act_owner,
                    target_date=to_datetime(treatment.target_date),
                    progress=treatment.progress or "0",
                    action_status_id=act_status,
                    next_followup_date=to_datetime(treatment.next_followup_date or treatment.target_date),
                    created_by=current_user.get("id", 1),
                    created_on=now_dt,
                    modified_by=current_user.get("id", 1),
                    modified_on=now_dt,
                    is_deleted=0
                )

                db.add(new_treatment)
                db.flush()

                saved_treatments.append(new_treatment)

                hist_treatment = RiskTreatmentHist(
                    risk_treatment_id=new_treatment.risk_treatment_id,
                    risk_description_id=new_treatment.risk_description_id,
                    risk_register_id=new_treatment.risk_register_id,
                    risk_id=new_treatment.risk_id,
                    action_plan=new_treatment.action_plan,
                    action_owner_id=new_treatment.action_owner_id,
                    target_date=new_treatment.target_date,
                    progress=new_treatment.progress,
                    action_status_id=new_treatment.action_status_id,
                    next_followup_date=new_treatment.next_followup_date,
                    created_by=new_treatment.created_by or current_user.get("id", 1),
                    created_on=new_treatment.created_on or now_dt,
                    modified_by=new_treatment.modified_by or current_user.get("id", 1),
                    modified_on=new_treatment.modified_on or now_dt,
                    is_deleted=new_treatment.is_deleted or 0
                )

                db.add(hist_treatment)

        db.commit()

        return {
            "risk_register": model_to_dict(risk),
            "risk_description": model_to_dict(description) if description else None,
            "risk_treatments": [model_to_dict(t) for t in saved_treatments]
        }

    except Exception as e:
        db.rollback()
        raise e



# Risk get by User (Optinal API)
def get_risk_by_user(db, user_id):

    # --- PREVIOUS UNOPTIMIZED CODE (Commented out) ---
    # risks = db.query(RiskRegister).filter(
    #     RiskRegister.risk_owner_id == user_id,
    #     RiskRegister.is_deleted == 0
    # ).all()
    # 
    # result = []
    # 
    # for risk in risks:
    # 
    #     description = db.query(RiskDescription).filter(
    #         RiskDescription.risk_register_id == risk.risk_register_id
    #     ).first()
    # 
    #     treatments = []
    # 
    #     if description:
    #         treatments = db.query(RiskTreatment).filter(
    #             RiskTreatment.risk_description_id == description.risk_description_id
    #         ).all()
    # 
    #     result.append({
    #         "risk_register_id": risk.risk_register_id,
    #         "risk_id": risk.risk_id,
    #         "risk_name": risk.risk_name,
    #         "financial_year": risk.financial_year,
    #         "risk_status": risk.risk_status,
    #         "risk_progress": risk.risk_progress,
    #         "description": description,
    #         "treatments": treatments
    #     })
    # 
    # return result

    # --- OPTIMIZED CODE (Eager loading relationships to eliminate nested query loops) ---
    risks = (
        db.query(RiskRegister)
        .options(
            joinedload(RiskRegister.risk_descriptions).joinedload(RiskDescription.risk_treatments)
        )
        .filter(
            RiskRegister.risk_owner_id == user_id,
            RiskRegister.is_deleted == 0
        )
        .all()
    )

    result = []
    for risk in risks:
        desc = risk.risk_descriptions[0] if risk.risk_descriptions else None
        treatments = desc.risk_treatments if desc and hasattr(desc, 'risk_treatments') else []
        result.append({
            "risk_register_id": risk.risk_register_id,
            "risk_id": risk.risk_id,
            "risk_name": risk.risk_name,
            "financial_year": risk.financial_year,
            "risk_status": risk.risk_status,
            "risk_progress": risk.risk_progress,
            "description": desc,
            "treatments": treatments
        })

    return result



# function use for converting SQLAlchemy model object to dict, if object is None then return dict with all keys but value as None, if prefix is provided then add prefix to all keys in the dict
def to_dict(obj, model=None,prefix=None):
    def format_key(key):
        return f"{prefix}{key}" if prefix else key
    
    # If object exists
    if obj is not None:
        return {
            format_key(c.key): getattr(obj, c.key)
            for c in inspect(obj).mapper.column_attrs
        }

    # If object is None → return all columns as None
    if model is not None:
        return {
            format_key(c.key): None
            for c in inspect(model).mapper.column_attrs
        }
    return {}


# find color code based on risk score, score is calculated by multiplying likelihood and impact, the color code is determined based on predefined thresholds for the score.
def get_color(score):
    
    if score <= 4:
        return "#4CAF50"
    elif score <= 9:
        return "#FFEB3B"
    elif score <= 16:
        return "#FF9800"
    else:
        return "#F44336"
        
COLOR_MAP = {
    "1A": "#008000",  # Green
    "2A": "#008000",
    "3A": "#008000",
    "1B": "#008000",
    "2B": "#008000",

    "4A": "#90EE90",  # Light Green
    "5A": "#90EE90",
    "3B": "#90EE90",
    "1C": "#90EE90",
    "2C": "#90EE90",

    "4B": "#FFFF00",  # Yellow
    "5B": "#FFFF00",
    "3C": "#FFFF00",
    "1D": "#FFFF00",
    "2D": "#FFFF00",

    "4C": "#FFA500",  # Orange
    "5C": "#FFA500",
    "3D": "#FFA500",
    "1E": "#FFA500",
    "2E": "#FFA500",

    "4D": "#FF0000",  # Red
    "5D": "#FF0000",
    "3E": "#FF0000",
    "4E": "#FF0000",
    "5E": "#FF0000",
}

def get_color_code(code: str) -> str:
    return COLOR_MAP.get(code.upper().strip(), "#FFFFFF")  # Default White


#-----------
# Get Risk by id (Assign id)
#----------
def get_risk_by_id(db, id):
    impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}
    try:
        # --- PREVIOUS UNOPTIMIZED CODE (Commented out) ---
        # query = db.query(
        #         RiskRegister,
        #         RiskDescription,
        #         RiskTreatment
        #     ).join(
        #         RiskDescription,
        #         RiskRegister.risk_register_id == RiskDescription.risk_register_id
        #     ).join(
        #         RiskTreatment,
        #         RiskDescription.risk_description_id == RiskTreatment.risk_description_id
        #     ).filter(
        #         RiskRegister.is_deleted == 0,
        #         RiskRegister.risk_function_head_approval_status == 1,
        #         RiskRegister.risk_head_approval_status == 1,
        #         RiskRegister.risk_manager_approval_status == 1
        #     )
        #     
        # if id:
        #     query = query.filter(RiskTreatment.action_owner_id == id)
        # 
        # records = query.order_by(
        #         RiskRegister.risk_register_id,
        #         RiskDescription.risk_description_id,
        #         RiskTreatment.risk_treatment_id
        #     ).all()

        # --- OPTIMIZED CODE (Eager loading relationships to eliminate lazy loading per row) ---
        query = (
            db.query(
                RiskRegister,
                RiskDescription,
                RiskTreatment
            )
            .options(
                joinedload(RiskRegister.risk_owner),
                joinedload(RiskRegister.risk_co_owner),
                joinedload(RiskTreatment.status)
            )
            .join(
                RiskDescription,
                RiskRegister.risk_register_id == RiskDescription.risk_register_id
            )
            .join(
                RiskTreatment,
                RiskDescription.risk_description_id == RiskTreatment.risk_description_id
            )
            .filter(
                RiskRegister.is_deleted == 0,
                RiskRegister.risk_function_head_approval_status == 1,
                RiskRegister.risk_head_approval_status == 1,
                RiskRegister.risk_manager_approval_status == 1
            )
        )
            
        if id:
            query = query.filter(RiskTreatment.action_owner_id == id)

        records = query.order_by(
                RiskRegister.risk_register_id,
                RiskDescription.risk_description_id,
                RiskTreatment.risk_treatment_id
            ).all()    
        

        #result = [ {**to_dict(rr), **to_dict(rd), **to_dict(rt)} for rr, rd, rt in records ]
        result = []

        for rr, rd, rt in records:

            # all_approved = (
            #     rr.function_head_status 
            #     and rr.function_head_status.status_name == "Approved"
            #     and rr.risk_head_status 
            #     and rr.risk_head_status.status_name == "Approved"
            #     and rr.risk_manager_status 
            #     and rr.risk_manager_status.status_name == "Approved"
            # )
            
            # all_approved = (
            #     rr.risk_function_head_approval_status == 1 and
            #     rr.risk_head_approval_status == 1 and
            #     rr.risk_manager_approval_status == 1
            # )

            # if not all_approved:
            #     continue

            risk_owner_name = rr.risk_owner.log_id if rr.risk_owner else None
            risk_co_owner_name = rr.risk_co_owner.log_id if rr.risk_co_owner else None
            risk_status_name = rt.status.status_name if rt.status else None

            likelihood = rd.inherent_risk_likelihood_id
            impact = rd.inherent_risk_impact_id
            current_likelihood = rd.current_risk_likelihood_id
            current_impact = rd.current_risk_impact_id

            inherent_color_str = None
            inherent_color_code = None
            current_color_str = None
            current_color_code = None

            if likelihood and impact:
                #inherent_color_str = get_color(likelihood*impact)
                inherent_color_code = f"{likelihood}{impact_map.get(impact)}"
                inherent_color_str = get_color_code(inherent_color_code)
                #current_color_str = get_color(current_likelihood*current_impact)
                current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"
                current_color_str = get_color_code(current_color_code)

            result.append({
                **to_dict(rr),
                **to_dict(rd),
                **to_dict(rt),
                "inherent_color_str": inherent_color_str,
                "inherent_color_code": inherent_color_code,
                "current_color_str": current_color_str,
                "current_color_code": current_color_code,
                "risk_owner_name": risk_owner_name,
                "risk_co_owner_name": risk_co_owner_name,
                "risk_status_name": risk_status_name
            })

        return result
    except Exception as e:
        raise e
    

def get_approval_status_name(status):
    
        status_map = {
            1: "Approved",
            -1: "Rejected",
            None: " "
        }

        return status_map.get(status, " ")

#-----------
# Risk LIST from Department id
#----------
# def get_risk_by_dept(db, dept_id,current_user):  
#     impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}
    
    
    
#     try:
#         dept_id_cur_user = current_user['dept_id']
#         user_type_name = current_user['user_type_name']
#         query = db.query(
#                     RiskRegister,
#                     RiskDescription
#                 ).outerjoin(
#                     RiskDescription,
#                     RiskRegister.risk_register_id == RiskDescription.risk_register_id
#                 ).join(
#                     Status,
#                     RiskRegister.risk_status == Status.id)
                
#         #query = query.filter(RiskRegister.is_deleted == 0)
        
#         # if user_type_name.upper() == 'RISK OWNER':
#         #     query = query.filter(RiskRegister.is_deleted == 0,RiskRegister.dept_id == dept_id_cur_user)
#         # elif user_type_name.upper() == 'FUNCTIONAL HEAD':
#         #     query = query.filter(RiskRegister.is_deleted == 0,RiskRegister.dept_id == dept_id_cur_user,
#         #                    func.upper(Status.status_name) != 'DRAFT' ,  
#         #                      func.upper(Status.status_name) != 'PENDING FOR ACTION'  )
#         # elif user_type_name.upper() == 'RISK MANAGER':
#         #     query = query.filter(RiskRegister.is_deleted == 0, RiskRegister.risk_function_head_approval_status == 1 and RiskRegister.risk_manager_approval_status != -1)
#         # elif user_type_name.upper() == 'RISK HEAD':
#         #     query = query.filter(RiskRegister.is_deleted == 0, RiskRegister.risk_manager_approval_status == 1 and RiskRegister.risk_head_approval_status != -1)
#         # else:
#         #     query = query.filter(RiskRegister.is_deleted == 0)
        
#         no_rejection_condition = and_(

#             or_(
#                 RiskRegister.risk_function_head_approval_status.is_(None),
#                 RiskRegister.risk_function_head_approval_status != -1
#             ),

#             or_(
#                 RiskRegister.risk_manager_approval_status.is_(None),
#                 RiskRegister.risk_manager_approval_status != -1
#             ),

#             or_(
#                 RiskRegister.risk_head_approval_status.is_(None),
#                 RiskRegister.risk_head_approval_status != -1
#             )
#         )
        
#         # ---------------- ROLE FILTER ----------------

#         if user_type_name.upper() == 'RISK OWNER':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,
#                 RiskRegister.dept_id == dept_id_cur_user
#             )

#         elif user_type_name.upper() == 'FUNCTIONAL HEAD':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,

#                 RiskRegister.dept_id == dept_id_cur_user,

#                 func.upper(Status.status_name) != 'DRAFT', func.upper(Status.status_name) != 'PENDING FOR ACTION',

#                 no_rejection_condition
#             )

#         elif user_type_name.upper() == 'RISK MANAGER':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,

#                 # FH approved
#                 RiskRegister.risk_function_head_approval_status == 1,

#                 no_rejection_condition
#             )

#         elif user_type_name.upper() == 'RISK HEAD':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,

#                 # RM approved
#                 RiskRegister.risk_manager_approval_status == 1,

#                 no_rejection_condition
#             )

#         else:

#             query = query.filter(
#                 RiskRegister.is_deleted == 0
#             )

#         if dept_id:
#             query = query.filter( RiskRegister.dept_id == dept_id )


#         records = query.order_by(
#                 RiskRegister.risk_register_id,
#                 RiskDescription.risk_description_id
#             ).all()    
        

#         result = []
#         for rr, rd in records:
#             risk_owner_name = None
#             if rr.risk_owner:
#                 risk_owner_name = rr.risk_owner.log_id
            
#             risk_status_name = None
#             if rr.status:
#                 risk_status_name = rr.status.status_name
                
#             # APPROVAL STATUS NAMES
#             # # read status from risk history table last entry
#             #risk_id_tmp = rr.risk_register_id
#             #risk_history = (
#             #    db.query(RiskRegisterHist)
#             #    .filter(RiskRegisterHist.risk_register_id == risk_id_tmp,
#             #            RiskRegisterHist.modified_on.isnot(None))
#             #    .order_by((RiskRegisterHist.modified_on).desc())
#             #    .first()
#             #)
            
#             function_head_approval_status_name = (
#                 get_approval_status_name(
#                     rr.risk_function_head_approval_status
#                 )
#             )

#             risk_manager_approval_status_name = (
#                 get_approval_status_name(
#                     rr.risk_manager_approval_status
#                 )
#             )

#             risk_head_approval_status_name = (
#                 get_approval_status_name(
#                     rr.risk_head_approval_status
#                 )
#             )

#             likelihood = None
#             impact = None
#             current_likelihood = None
#             current_impact = None
#             inherent_color_str = None
#             inherent_color_code = None
#             current_color_str = None
#             current_color_code = None
           
#             if rd is not None:
#                 if rd.inherent_risk_likelihood_id:
#                     likelihood = rd.inherent_risk_likelihood_id
#                 if rd.inherent_risk_impact_id:
#                     impact = rd.inherent_risk_impact_id

#                 if rd.current_risk_likelihood_id:
#                     current_likelihood = rd.current_risk_likelihood_id
#                 if rd.current_risk_impact_id:
#                     current_impact = rd.current_risk_impact_id


#                 if likelihood and impact:
#                     #inherent_color_str = get_color(likelihood*impact)
#                     inherent_color_code = f"{likelihood}{impact_map.get(impact)}"
#                     inherent_color_str = get_color_code(inherent_color_code)

#                 if current_likelihood and current_impact:
#                     #current_color_str = get_color(current_likelihood*current_impact)
#                     current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"
#                     current_color_str = get_color_code(current_color_code)

#             result.append({
#                 **to_dict(rr),
#                 **to_dict(rd, RiskDescription, prefix="rd_"),
#                 "inherent_color_str": inherent_color_str,
#                 "inherent_color_code" : inherent_color_code,
#                 "current_color_str": current_color_str,
#                 "current_color_code" : current_color_code,
#                 "risk_owner_name" : risk_owner_name,
#                 "risk_status_name" : risk_status_name,
                
#                 "function_head_approval_status_name":
#                     function_head_approval_status_name,

#                 "risk_manager_approval_status_name":
#                     risk_manager_approval_status_name,

#                 "risk_head_approval_status_name":
#                     risk_head_approval_status_name
#             })

#         return result
#     except Exception as e:
#         raise e

#----------------------------------------------------------new function------------

# def get_risk_by_dept(db, dept_id,financial_year,current_user):  
#     impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}
    
    
    
#     try:
#         dept_id_cur_user = current_user['dept_id']
#         user_type_name = current_user['user_type_name']
#         query = db.query(
#                     RiskRegister,
#                     RiskDescription
#                 ).outerjoin(
#                     RiskDescription,
#                     RiskRegister.risk_register_id == RiskDescription.risk_register_id
#                 ).join(
#                     Status,
#                     RiskRegister.risk_status == Status.id)
                
#         #query = query.filter(RiskRegister.is_deleted == 0)
        
#         # if user_type_name.upper() == 'RISK OWNER':
#         #     query = query.filter(RiskRegister.is_deleted == 0,RiskRegister.dept_id == dept_id_cur_user)
#         # elif user_type_name.upper() == 'FUNCTIONAL HEAD':
#         #     query = query.filter(RiskRegister.is_deleted == 0,RiskRegister.dept_id == dept_id_cur_user,
#         #                    func.upper(Status.status_name) != 'DRAFT' ,  
#         #                      func.upper(Status.status_name) != 'PENDING FOR ACTION'  )
#         # elif user_type_name.upper() == 'RISK MANAGER':
#         #     query = query.filter(RiskRegister.is_deleted == 0, RiskRegister.risk_function_head_approval_status == 1 and RiskRegister.risk_manager_approval_status != -1)
#         # elif user_type_name.upper() == 'RISK HEAD':
#         #     query = query.filter(RiskRegister.is_deleted == 0, RiskRegister.risk_manager_approval_status == 1 and RiskRegister.risk_head_approval_status != -1)
#         # else:
#         #     query = query.filter(RiskRegister.is_deleted == 0)
        
#         no_rejection_condition = and_(

#             or_(
#                 RiskRegister.risk_function_head_approval_status.is_(None),
#                 RiskRegister.risk_function_head_approval_status != -1
#             ),

#             or_(
#                 RiskRegister.risk_manager_approval_status.is_(None),
#                 RiskRegister.risk_manager_approval_status != -1
#             ),

#             or_(
#                 RiskRegister.risk_head_approval_status.is_(None),
#                 RiskRegister.risk_head_approval_status != -1
#             )
#         )
        
#         # ---------------- ROLE FILTER ----------------

#         if user_type_name.upper() == 'RISK OWNER':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,
#                 RiskRegister.dept_id == dept_id_cur_user
#             )

#         elif user_type_name.upper() == 'FUNCTIONAL HEAD':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,

#                 RiskRegister.dept_id == dept_id_cur_user,

#                 func.upper(Status.status_name) != 'DRAFT', func.upper(Status.status_name) != 'PENDING FOR ACTION',

#                 no_rejection_condition
#             )

#         elif user_type_name.upper() == 'RISK MANAGER':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,

#                 # FH approved
#                 RiskRegister.risk_function_head_approval_status == 1,

#                 no_rejection_condition
#             )

#         elif user_type_name.upper() == 'RISK HEAD':

#             query = query.filter(
#                 RiskRegister.is_deleted == 0,

#                 # RM approved
#                 RiskRegister.risk_manager_approval_status == 1,

#                 no_rejection_condition
#             )

#         else:

#             query = query.filter(
#                 RiskRegister.is_deleted == 0
#             )

#         if dept_id:
#             query = query.filter( RiskRegister.dept_id == dept_id )
            
            
#         # Financial Year Filter. if not pass give current only
#         if financial_year:
#             query = query.filter(
#                 RiskRegister.financial_year == financial_year
#             )
#         else:
#             # Current financial year
#             today = datetime.now()

#             if today.month >= 4:
#                 current_fy = f"{today.year}-{today.year + 1}"
#             else:
#                 current_fy = f"{today.year - 1}-{today.year}"

#             query = query.filter(
#                 RiskRegister.financial_year == current_fy
#     )


#         records = query.order_by(
#                 RiskRegister.risk_register_id,
#                 RiskDescription.risk_description_id
#             ).all()    
        
#         action_required = False
#         result = []
#         for rr, rd in records:
#             action_required = False
#             if rd:
#                 action_required = ( 
#                     db.query(RiskTreatment)
#                     .filter(RiskTreatment.risk_register_id == rr.risk_register_id,
#                             RiskTreatment.progress == '100',
                            
#                             or_(
#                                 RiskTreatment.approval_status.notin_([-1, 1]),
#                                 RiskTreatment.approval_status.is_(None)
#                             )
#                             ).first()
#                             is not None
#                 )
             
#             risk_owner_name = None
#             if rr.risk_owner:
#                 risk_owner_name = rr.risk_owner.log_id
            
#             risk_status_name = None
#             if rr.status:
#                 risk_status_name = rr.status.status_name
                
#             # APPROVAL STATUS NAMES
#             # # read status from risk history table last entry
#             #risk_id_tmp = rr.risk_register_id
#             #risk_history = (
#             #    db.query(RiskRegisterHist)
#             #    .filter(RiskRegisterHist.risk_register_id == risk_id_tmp,
#             #            RiskRegisterHist.modified_on.isnot(None))
#             #    .order_by((RiskRegisterHist.modified_on).desc())
#             #    .first()
#             #)
            
#             function_head_approval_status_name = (
#                 get_approval_status_name(
#                     rr.risk_function_head_approval_status
#                 )
#             )

#             risk_manager_approval_status_name = (
#                 get_approval_status_name(
#                     rr.risk_manager_approval_status
#                 )
#             )

#             risk_head_approval_status_name = (
#                 get_approval_status_name(
#                     rr.risk_head_approval_status
#                 )
#             )

#             likelihood = None
#             impact = None
#             current_likelihood = None
#             current_impact = None
#             inherent_color_str = None
#             inherent_color_code = None
#             current_color_str = None
#             current_color_code = None
           
#             if rd is not None:
#                 if rd.inherent_risk_likelihood_id:
#                     likelihood = rd.inherent_risk_likelihood_id
#                 if rd.inherent_risk_impact_id:
#                     impact = rd.inherent_risk_impact_id

#                 if rd.current_risk_likelihood_id:
#                     current_likelihood = rd.current_risk_likelihood_id
#                 if rd.current_risk_impact_id:
#                     current_impact = rd.current_risk_impact_id


#                 if likelihood and impact:
#                     #inherent_color_str = get_color(likelihood*impact)
#                     inherent_color_code = f"{likelihood}{impact_map.get(impact)}"
#                     inherent_color_str = get_color_code(inherent_color_code)

#                 if current_likelihood and current_impact:
#                     #current_color_str = get_color(current_likelihood*current_impact)
#                     current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"
#                     current_color_str = get_color_code(current_color_code)

#             result.append({
#                 **to_dict(rr),
#                 **to_dict(rd, RiskDescription, prefix="rd_"),
#                 "inherent_color_str": inherent_color_str,
#                 "inherent_color_code" : inherent_color_code,
#                 "current_color_str": current_color_str,
#                 "current_color_code" : current_color_code,
#                 "risk_owner_name" : risk_owner_name,
#                 "risk_status_name" : risk_status_name,
                
#                 "function_head_approval_status_name":
#                     function_head_approval_status_name,

#                 "risk_manager_approval_status_name":
#                     risk_manager_approval_status_name,

#                 "risk_head_approval_status_name":
#                     risk_head_approval_status_name,
#                 "action_required": action_required
#             })

#         return result
#     except Exception as e:
#         raise e
    
    
    
    
#---------------------new function with no restriction in FH, RM, RH---------------------
def get_risk_by_dept(db, dept_id,financial_year,current_user):  
    impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}
    
    
    
    try:
        dept_id_cur_user = current_user['dept_id']
        user_type_name = current_user['user_type_name']
        # --- PREVIOUS UNOPTIMIZED CODE (Commented out) ---
        # query = db.query(
        #             RiskRegister,
        #             RiskDescription
        #         ).outerjoin(
        #             RiskDescription,
        #             RiskRegister.risk_register_id == RiskDescription.risk_register_id
        #         ).join(
        #             Status,
        #             RiskRegister.risk_status == Status.id)

        # --- OPTIMIZED CODE (Eager load relationships to eliminate N+1) ---
        query = (
            db.query(
                RiskRegister,
                RiskDescription
            )
            .options(
                joinedload(RiskRegister.risk_owner),
                joinedload(RiskRegister.status)
            )
            .outerjoin(
                RiskDescription,
                RiskRegister.risk_register_id == RiskDescription.risk_register_id
            )
            .join(
                Status,
                RiskRegister.risk_status == Status.id
            )
        )
                
        #query = query.filter(RiskRegister.is_deleted == 0)
        
        # if user_type_name.upper() == 'RISK OWNER':
        #     query = query.filter(RiskRegister.is_deleted == 0,RiskRegister.dept_id == dept_id_cur_user)
        # elif user_type_name.upper() == 'FUNCTIONAL HEAD':
        #     query = query.filter(RiskRegister.is_deleted == 0,RiskRegister.dept_id == dept_id_cur_user,
        #                    func.upper(Status.status_name) != 'DRAFT' ,  
        #                      func.upper(Status.status_name) != 'PENDING FOR ACTION'  )
        # elif user_type_name.upper() == 'RISK MANAGER':
        #     query = query.filter(RiskRegister.is_deleted == 0, RiskRegister.risk_function_head_approval_status == 1 and RiskRegister.risk_manager_approval_status != -1)
        # elif user_type_name.upper() == 'RISK HEAD':
        #     query = query.filter(RiskRegister.is_deleted == 0, RiskRegister.risk_manager_approval_status == 1 and RiskRegister.risk_head_approval_status != -1)
        # else:
        #     query = query.filter(RiskRegister.is_deleted == 0)
        
        no_rejection_condition = and_(

            or_(
                RiskRegister.risk_function_head_approval_status.is_(None),
                RiskRegister.risk_function_head_approval_status != -1
            ),

            or_(
                RiskRegister.risk_manager_approval_status.is_(None),
                RiskRegister.risk_manager_approval_status != -1
            ),

            or_(
                RiskRegister.risk_head_approval_status.is_(None),
                RiskRegister.risk_head_approval_status != -1
            )
        )
        
        # ---------------- ROLE FILTER ----------------

        # if user_type_name.upper() == 'RISK OWNER':

        #     query = query.filter(
        #         RiskRegister.is_deleted == 0,
        #         RiskRegister.dept_id == dept_id_cur_user
        #     )
        
        if user_type_name.upper() == 'RISK OWNER':

            query = query.filter(
                RiskRegister.is_deleted == 0,
                or_(
                    # Risk Owner: same department
                    and_(
                        RiskRegister.dept_id == dept_id_cur_user,
                        RiskRegister.risk_owner_id == current_user["id"]
                    ),

                    # Risk Co-Owner: any department
                    RiskRegister.risk_co_owner_id == current_user["id"]
                )
            )

        elif user_type_name.upper() == 'FUNCTIONAL HEAD':

            query = query.filter(
                RiskRegister.is_deleted == 0,
                
                or_(
            RiskRegister.dept_id == dept_id_cur_user,
            RiskRegister.risk_co_owner_id == current_user["id"]
        ),

                func.upper(Status.status_name) != 'DRAFT', func.upper(Status.status_name) != 'PENDING FOR ACTION',

                no_rejection_condition
            )

        elif user_type_name.upper() == 'RISK MANAGER':

            query = query.filter(
                RiskRegister.is_deleted == 0,

                # FH approved
                #RiskRegister.risk_function_head_approval_status == 1,
                func.upper(Status.status_name) != 'DRAFT', func.upper(Status.status_name) != 'PENDING FOR ACTION',

                no_rejection_condition
            )

        elif user_type_name.upper() == 'RISK HEAD':

            query = query.filter(
                RiskRegister.is_deleted == 0,

                # RM approved
                #RiskRegister.risk_manager_approval_status == 1,
                func.upper(Status.status_name) != 'DRAFT', func.upper(Status.status_name) != 'PENDING FOR ACTION',

                no_rejection_condition
            )

        else:

            query = query.filter(
                RiskRegister.is_deleted == 0
            )

        if dept_id:
            query = query.filter( RiskRegister.dept_id == dept_id )

        if financial_year:
            query = query.filter(RiskRegister.financial_year == financial_year)
        else:
            today = datetime.now()
            if today.month >= 4:
                current_fy = f"{today.year}-{today.year + 1}"
            else:
                current_fy = f"{today.year - 1}-{today.year}"
                
            query = query.filter(RiskRegister.financial_year == current_fy)


        records = query.order_by(
                RiskRegister.risk_register_id,
                RiskDescription.risk_description_id
            ).all()    
        
        # --- PREVIOUS UNOPTIMIZED CODE (Commented out) ---
        # action_required = False
        # result = []
        # for rr, rd in records:
        #     action_required = False
        #     if rd:
        #         action_required = ( 
        #             db.query(RiskTreatment)
        #             .filter(RiskTreatment.risk_register_id == rr.risk_register_id,
        #                     RiskTreatment.progress == '100',
        #                     
        #                     or_(
        #                         RiskTreatment.approval_status.notin_([-1, 1]),
        #                         RiskTreatment.approval_status.is_(None)
        #                     )
        #                     ).first()
        #                     is not None
        #         )
        #      
        #     risk_owner_name = None
        #     if rr.risk_owner:
        #         risk_owner_name = rr.risk_owner.log_id
        #     
        #     risk_status_name = None
        #     if rr.status:
        #         risk_status_name = rr.status.status_name

        # --- OPTIMIZED CODE (Batch fetch action_required IDs in 1 single query instead of querying in a loop) ---
        risk_register_ids = [rr.risk_register_id for rr, rd in records if rd]
        action_required_risk_ids = set()
        if risk_register_ids:
            action_treatments = (
                db.query(RiskTreatment.risk_register_id)
                .filter(
                    RiskTreatment.risk_register_id.in_(risk_register_ids),
                    RiskTreatment.progress == '100',
                    or_(
                        RiskTreatment.approval_status.notin_([-1, 1]),
                        RiskTreatment.approval_status.is_(None)
                    )
                )
                .all()
            )
            action_required_risk_ids = {row[0] for row in action_treatments}

        result = []
        for rr, rd in records:
            action_required = rr.risk_register_id in action_required_risk_ids if rd else False
            risk_owner_name = rr.risk_owner.log_id if rr.risk_owner else None
            risk_status_name = rr.status.status_name if rr.status else None
                
            # APPROVAL STATUS NAMES
            # # read status from risk history table last entry
            #risk_id_tmp = rr.risk_register_id
            #risk_history = (
            #    db.query(RiskRegisterHist)
            #    .filter(RiskRegisterHist.risk_register_id == risk_id_tmp,
            #            RiskRegisterHist.modified_on.isnot(None))
            #    .order_by((RiskRegisterHist.modified_on).desc())
            #    .first()
            #)
            
            function_head_approval_status_name = (
                get_approval_status_name(
                    rr.risk_function_head_approval_status
                )
            )

            risk_manager_approval_status_name = (
                get_approval_status_name(
                    rr.risk_manager_approval_status
                )
            )

            risk_head_approval_status_name = (
                get_approval_status_name(
                    rr.risk_head_approval_status
                )
            )

            likelihood = None
            impact = None
            current_likelihood = None
            current_impact = None
            inherent_color_str = None
            inherent_color_code = None
            current_color_str = None
            current_color_code = None
           
            if rd is not None:
                if rd.inherent_risk_likelihood_id:
                    likelihood = rd.inherent_risk_likelihood_id
                if rd.inherent_risk_impact_id:
                    impact = rd.inherent_risk_impact_id

                if rd.current_risk_likelihood_id:
                    current_likelihood = rd.current_risk_likelihood_id
                if rd.current_risk_impact_id:
                    current_impact = rd.current_risk_impact_id


                if likelihood and impact:
                    #inherent_color_str = get_color(likelihood*impact)
                    inherent_color_code = f"{likelihood}{impact_map.get(impact)}"
                    inherent_color_str = get_color_code(inherent_color_code)

                if current_likelihood and current_impact:
                    #current_color_str = get_color(current_likelihood*current_impact)
                    current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"
                    current_color_str = get_color_code(current_color_code)

            result.append({
                **to_dict(rr),
                **to_dict(rd, RiskDescription, prefix="rd_"),
                "inherent_color_str": inherent_color_str,
                "inherent_color_code" : inherent_color_code,
                "current_color_str": current_color_str,
                "current_color_code" : current_color_code,
                "risk_owner_name" : risk_owner_name,
                "risk_status_name" : risk_status_name,
                
                "function_head_approval_status_name":
                    function_head_approval_status_name,

                "risk_manager_approval_status_name":
                    risk_manager_approval_status_name,

                "risk_head_approval_status_name":
                    risk_head_approval_status_name,
                "action_required": action_required
            })

        return result
    except Exception as e:
        raise e


#-----------
# Risk Register by risk_id
#------------

def get_risk_by_risk_id(db, risk_id):
    impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}

    try:
        risks = (
            db.query(RiskRegister)
            .options(
                joinedload(RiskRegister.risk_owner),
                joinedload(RiskRegister.risk_co_owner),
                joinedload(RiskRegister.status),

                joinedload(RiskRegister.function_head_status),
                joinedload(RiskRegister.risk_head_status),
                joinedload(RiskRegister.risk_manager_status),

                joinedload(RiskRegister.risk_function_head_approval_by_name),
                joinedload(RiskRegister.risk_head_approval_by_name),
                joinedload(RiskRegister.risk_manager_approval_by_name),

                joinedload(RiskRegister.risk_descriptions)
                .joinedload(RiskDescription.treatments)
                .joinedload(RiskTreatment.action_owner),

                joinedload(RiskRegister.risk_descriptions)
                .joinedload(RiskDescription.treatments)
                .joinedload(RiskTreatment.status)
            )
            .filter(
                RiskRegister.risk_register_id == risk_id,
                RiskRegister.is_deleted == 0
            )
            .all()
        )

        result = []

        for rr in risks:

            risk_dict = to_dict(rr)

            # ---------- Owner & Status ----------
            risk_dict["rd_risk_owner_name"] = rr.risk_owner.log_id if rr.risk_owner else None
            risk_dict["rd_risk_co_owner_name"] = rr.risk_co_owner.log_id if rr.risk_co_owner else None
            risk_dict["rd_risk_status_name"] = rr.status.status_name if rr.status else None

            # ---------- Function Head ----------
            if rr.risk_function_head_approval_status == 1:
                fh_status = "Approved"
            elif rr.risk_function_head_approval_status == -1:
                fh_status = "Rejected"
            else:
                fh_status = None

            risk_dict["risk_function_head_approval_status_name"] = fh_status
            
            
            risk_dict["risk_function_head_approval_by_name"] = (
                rr.risk_function_head_approval_by_name.log_id if rr.risk_function_head_approval_by_name else None
            )

            # ---------- Risk Head ----------
            if rr.risk_head_approval_status == 1:
                rh_status = "Approved"
            elif rr.risk_head_approval_status == -1:
                rh_status = "Rejected"
            else:
                rh_status = None

            risk_dict["risk_head_approval_status_name"] = rh_status
            
            
            risk_dict["risk_head_approval_by_name"] = (
                rr.risk_head_approval_by_name.log_id if rr.risk_head_approval_by_name else None
            )

            # ---------- Risk Manager ----------
            if rr.risk_manager_approval_status == 1:
                rm_status = "Approved"
            elif rr.risk_manager_approval_status == -1:
                rm_status = "Rejected"
            else:
                rm_status = None

            risk_dict["risk_manager_approval_status_name"] = rm_status
            
            
            risk_dict["risk_manager_approval_by_name"] = (
                rr.risk_manager_approval_by_name.log_id if rr.risk_manager_approval_by_name else None
            )

            # ---------- Final Approval Status ----------
            functional_head = rr.risk_function_head_approval_status
            risk_head = rr.risk_head_approval_status
            risk_manager = rr.risk_manager_approval_status

            if  rr.status and rr.status.status_name == 'Draft' or rr.status and rr.status.status_name == 'Pending for Action':
                final_status = rr.status.status_name
            else:
                if functional_head == 1 and risk_head == 1 and risk_manager == 1:
                    final_status = "Approved"

                elif functional_head == -1 or risk_head == -1 or risk_manager == -1:
                    final_status = "Rejected"

                else:
                    final_status = "Pending for Approval"

            risk_dict["approval_status"] = final_status
            
            # ---------- Risk Descriptions ----------
            risk_desc_list = []

            # for rd in rr.risk_descriptions:
            for rd in sorted(
                rr.risk_descriptions,
                key=lambda x: x.risk_description_id
            ):

                likelihood = rd.inherent_risk_likelihood_id
                impact = rd.inherent_risk_impact_id
                current_likelihood = rd.current_risk_likelihood_id
                current_impact = rd.current_risk_impact_id

                inherent_color_str = None
                inherent_color_code = None
                current_color_str = None
                current_color_code = None

                if likelihood and impact:
                    #inherent_color_str = get_color(likelihood * impact)
                    inherent_color_code = f"{likelihood}{impact_map.get(impact)}"
                    inherent_color_str = get_color_code(inherent_color_code)

                if current_likelihood and current_impact:
                    #current_color_str = get_color(current_likelihood * current_impact)
                    current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"
                    current_color_str = get_color_code(current_color_code)


                # ---------- Treatments ----------
                treatments_list = []

                # for rt in rd.treatments:
                for rt in sorted(
                    rd.treatments,
                    key=lambda x: x.risk_treatment_id
                ):
                    treatments_list.append({
                        **to_dict(rt),
                        "risk_owner_name": rt.action_owner.log_id if rt.action_owner else None,
                        "risk_co_owner_name": rt.action_owner.log_id if rt.action_owner else None,
                        "risk_status_name": rt.status.status_name if rt.status else None
                    })

                rd_dict = {
                    **to_dict(rd),
                    "inherent_color_str": inherent_color_str,
                    "inherent_color_code": inherent_color_code,
                    "current_color_str": current_color_str,
                    "current_color_code": current_color_code,
                    "treatments": treatments_list
                }

                risk_desc_list.append(rd_dict)

            risk_dict["risk_descriptions"] = risk_desc_list
            result.append(risk_dict)

        return result

    except Exception as e:
        raise e
    
    
# Risk Register by risk_description_id
# def get_risk_by_description_id(db, description_id):
#     try:

#         risk_description = (
#             db.query(RiskDescription)
#             .options(
                
#                 joinedload(RiskDescription.treatments)
#                 .joinedload(RiskTreatment.status)
#                 .load_only(Status.status_name),

#                 joinedload(RiskDescription.treatments)
#                 .joinedload(RiskTreatment.action_owner)
#                 .load_only(User.log_id)
#                 )
#             .filter(
#                 RiskDescription.risk_description_id == description_id,
#                 RiskDescription.is_deleted == 0
#             )
#             .first()
#         )

#         result = risk_description.__dict__.copy()

#         treatments_list = []

#         for t in risk_description.treatments:
#             treatment_dict = t.__dict__.copy()

#             # remove SQLAlchemy internal state
#             treatment_dict.pop("_sa_instance_state", None)

#             # add status_name
#             treatment_dict["risk_status_name"] = (
#                 t.status.status_name if t.status else None
#             )
#             treatment_dict["risk_owner_name"] = (
#                 t.action_owner.log_id if t.action_owner else None
#             )

#             # remove nested status object
#             treatment_dict.pop("status", None)
#             treatment_dict.pop("action_owner", None)

#             treatments_list.append(treatment_dict)

#         result["treatments"] = treatments_list
#         result.pop("_sa_instance_state", None)
#         return result

#     except Exception as e:
#         raise e


#----------
# Risk Register by risk_description_id
#----------

def get_risk_by_description_id(db, description_id):

    impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}

    try:

        descriptions = (
            db.query(RiskDescription)
            .options(
                joinedload(RiskDescription.treatments)
                .joinedload(RiskTreatment.action_owner),

                joinedload(RiskDescription.treatments)
                .joinedload(RiskTreatment.status),

                joinedload(RiskDescription.risk_register)
            )
            .filter(
                RiskDescription.risk_description_id == description_id,
                RiskDescription.is_deleted == 0
            )
            .all()
        )

        result = []

        # for rd in descriptions:
        for rd in sorted(
            descriptions,
            key=lambda x: x.risk_description_id
        ):

            likelihood = rd.inherent_risk_likelihood_id
            impact = rd.inherent_risk_impact_id
            current_likelihood = rd.current_risk_likelihood_id
            current_impact = rd.current_risk_impact_id

            inherent_color_str = None
            inherent_color_code = None
            current_color_str = None
            current_color_code = None

            if likelihood and impact:
                #inherent_color_str = get_color(likelihood * impact)
                inherent_color_code = f"{likelihood}{impact_map.get(impact)}"
                inherent_color_str = get_color_code(inherent_color_code)

            if current_likelihood and current_impact:
                #current_color_str = get_color(current_likelihood * current_impact)
                current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"
                current_color_str = get_color_code(current_color_code)

            # Treatments
            treatments_list = []

            # for rt in rd.treatments:
            for rt in sorted(
                rd.treatments,
                key=lambda x: x.risk_treatment_id
            ):
            
                treatments_list.append({
                    **to_dict(rt),
                    "action_owner_name": (
                        f"{rt.action_owner.first_name or ''} "
                        f"{rt.action_owner.last_name or ''}".strip()
                        if rt.action_owner else None
                    ),
                    #"action_owner_name": rt.action_owner.log_id if rt.action_owner else None,
                    #"risk_co_owner_name": rt.action_owner.log_id if rt.action_owner else None,
                    "risk_status_name": rt.status.status_name if rt.status else None
                })

            rd_dict = {
                **to_dict(rd),
                "inherent_color_str": inherent_color_str,
                "inherent_color_code": inherent_color_code,
                "current_color_str": current_color_str,
                "current_color_code": current_color_code,
                "treatments": treatments_list
            }

            # Risk Register info
            risk_dict = {
                **to_dict(rd.risk_register),
                "risk_owner_name": (
                    f"{rd.risk_register.risk_owner.first_name or ''} "
                    f"{rd.risk_register.risk_owner.last_name or ''}".strip()
                    if rd.risk_register.risk_owner else None
                ),
                "risk_co_owner_name": (
                    f"{rd.risk_register.risk_co_owner.first_name or ''} "
                    f"{rd.risk_register.risk_co_owner.last_name or ''}".strip()
                    if rd.risk_register.risk_co_owner else None
                ),
                "risk_descriptions": [rd_dict]
            }

            result.append(risk_dict)

        return result

    except Exception as e:
        raise e
    
    
#---------    
# Download Risk data in Excel format
#----------

def calculate_row_height(text, column_width):
    if not text:
        return 15
    
    # Estimate characters per line
    chars_per_line = column_width * 1.2
    
    # Estimate number of lines
    lines = math.ceil(len(str(text)) / chars_per_line)
    
    # Excel default row height ≈ 15
    return max(15, lines * 15)

    
# def get_risk_data_excel(db, dept_id):
    
#     impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}

#     green_value = {"1A", "2A", "3A", "1B", "2B"}
#     light_green_value = {"1C","2C", "3B","4A","5A"}
#     yellow_value = {"4B", "5B", "3C", "2D","1D"}
#     amber_value = {"1E", "2E", "3D", "4C", "5C"}
#     red_value = {"3E","4E", "5E", "4D", "5D"}

#     green_fill = PatternFill(start_color="008000", end_color="008000", fill_type="solid")
#     yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#     amber_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
#     red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
#     light_green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

#     try:

#         query = db.query(RiskRegister).options(
#             joinedload(RiskRegister.risk_descriptions)
#             .joinedload(RiskDescription.treatments)
#             .joinedload(RiskTreatment.action_owner),
#             joinedload(RiskRegister.department),
#             joinedload(RiskRegister.risk_owner)
#         )

#         if dept_id:
#             query = query.filter(RiskRegister.dept_id == dept_id)

#         query = query.filter(
#             RiskRegister.is_deleted == 0,
#             RiskRegister.dept_id > 0
#         ).order_by(RiskRegister.risk_register_id)

#         risks = query.all()

#         department_rows = {}
#         risk_merge_ranges = {}
#         description_merge_ranges = {}
#         current_rows = {}

#         for risk in risks:

#             dept_name = "UNKNOWN"
#             if risk.department:
#                 dept_name = risk.department.dept_short_name

#             if dept_name not in department_rows:
#                 department_rows[dept_name] = []
#                 risk_merge_ranges[dept_name] = []
#                 description_merge_ranges[dept_name] = []
#                 current_rows[dept_name] = 2

#             risk_owner_name = risk.risk_owner.log_id if risk.risk_owner else ""

#             risk_start_row = current_rows[dept_name]

#             descriptions = risk.risk_descriptions if risk.risk_descriptions else [None]

#             for desc in descriptions:

#                 desc_start_row = current_rows[dept_name]

#                 likelihood = desc.inherent_risk_likelihood_id if desc else None
#                 impact = desc.inherent_risk_impact_id if desc else None

#                 current_likelihood = desc.current_risk_likelihood_id if desc else None
#                 current_impact = desc.current_risk_impact_id if desc else None

#                 inherent_color_code = ""
#                 current_color_code = ""

#                 if likelihood and impact:
#                     inherent_color_code = f"{likelihood}{impact_map.get(impact)}"

#                 if current_likelihood and current_impact:
#                     current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"

#                 treatments = desc.treatments if desc and desc.treatments else [None]

#                 first_desc = True

#                 for treatment in treatments:

#                     department_rows[dept_name].append({

#                         "Risk ID": risk.risk_id if current_rows[dept_name] == risk_start_row else "",
#                         "Risk Name": risk.risk_name if current_rows[dept_name] == risk_start_row else "",

#                         "Risk Description": desc.risk_description if desc and first_desc else "",
#                         "Inherent Risk Level": inherent_color_code if first_desc else "",
#                         "Current Mitigation": desc.mitigation if desc and first_desc else "",
#                         "Current Risk Level": current_color_code if first_desc else "",

#                         "Risk Owner": risk_owner_name if current_rows[dept_name] == risk_start_row else "",
#                         "Action Plan": treatment.action_plan if treatment else "",
#                         "Action Owner": treatment.action_owner.log_id if treatment and treatment.action_owner else "",
#                         "Target Date": treatment.target_date.date() if treatment and treatment.target_date else "",
#                         "Action Status": treatment.status.status_name if treatment and treatment.status else ""
#                     })

#                     first_desc = False
#                     current_rows[dept_name] += 1

#                 desc_end_row = current_rows[dept_name] - 1

#                 if desc_end_row >= desc_start_row:
#                     description_merge_ranges[dept_name].append(
#                         (desc_start_row, desc_end_row)
#                     )

#             risk_end_row = current_rows[dept_name] - 1

#             if risk_end_row >= risk_start_row:
#                 risk_merge_ranges[dept_name].append(
#                     (risk_start_row, risk_end_row)
#                 )

#         output = io.BytesIO()

#         wrap_alignment = Alignment(wrap_text=True, vertical="center")

#         merge_alignment = Alignment(
#             horizontal="center",
#             vertical="center",
#             wrap_text=True
#         )

#         with pd.ExcelWriter(output, engine="openpyxl") as writer:

#             for dept_name, rows in department_rows.items():

#                 df = pd.DataFrame(rows)

#                 sheet_name = dept_name[:31]

#                 df.to_excel(writer, index=False, sheet_name=sheet_name)

#                 worksheet = writer.sheets[sheet_name]

#                 # Color Fill
#                 for row in range(2, len(df) + 2):

#                     cell_d = worksheet[f"D{row}"]
#                     cell_f = worksheet[f"F{row}"]

#                     if cell_d.value in green_value:
#                         cell_d.fill = green_fill
#                     elif cell_d.value in amber_value:
#                         cell_d.fill = amber_fill
#                     elif cell_d.value in yellow_value:
#                         cell_d.fill = yellow_fill
#                     elif cell_d.value in red_value:
#                         cell_d.fill = red_fill
#                     elif cell_d.value in light_green_value:
#                         cell_d.fill = light_green_fill

#                     if cell_f.value in green_value:
#                         cell_f.fill = green_fill
#                     elif cell_f.value in amber_value:
#                         cell_f.fill = amber_fill
#                     elif cell_f.value in yellow_value:
#                         cell_f.fill = yellow_fill
#                     elif cell_f.value in red_value:
#                         cell_f.fill = red_fill
#                     elif cell_f.value in light_green_value:
#                         cell_f.fill = light_green_fill

#                 # Column Width
#                 column_widths = {
#                     "A": 15,   # Risk ID (A)
#                     "B" :20, # Risk Name (B)
#                     "C": 50,   # Risk Description (C)
#                     "D": 10,   # Inherent Risk Level (D)
#                     "E": 70,    # Mitigation (E)
#                     "F": 10, # Current risk level (F)
#                     "G": 15, # Risk Owner (G)
#                     "H": 70, # Treatment (H)
#                     "I": 15, # Action Owner (I)
#                     "J": 15,  # Target Date (J)
#                     "K":15 # Status
#                 }

#                 for col, width in column_widths.items():
#                     worksheet.column_dimensions[col].width = width

#                 # Risk Merge
#                 for start, end in risk_merge_ranges[dept_name]:

#                     if start != end:

#                         worksheet.merge_cells(f"A{start}:A{end}")
#                         worksheet.merge_cells(f"B{start}:B{end}")
#                         worksheet.merge_cells(f"G{start}:G{end}")

#                     worksheet[f"A{start}"].alignment = merge_alignment
#                     worksheet[f"B{start}"].alignment = merge_alignment
#                     worksheet[f"G{start}"].alignment = merge_alignment

#                 # Description Merge
#                 for start, end in description_merge_ranges[dept_name]:

#                     if start != end:

#                         worksheet.merge_cells(f"C{start}:C{end}")
#                         worksheet.merge_cells(f"D{start}:D{end}")
#                         worksheet.merge_cells(f"E{start}:E{end}")
#                         worksheet.merge_cells(f"F{start}:F{end}")

#                     worksheet[f"C{start}"].alignment = merge_alignment
#                     worksheet[f"D{start}"].alignment = merge_alignment
#                     worksheet[f"E{start}"].alignment = merge_alignment
#                     worksheet[f"F{start}"].alignment = merge_alignment

#                 # Wrap Text & Dynamic Height
#                 for row in range(2, len(df) + 2):

#                     worksheet[f"C{row}"].alignment = wrap_alignment
#                     worksheet[f"E{row}"].alignment = wrap_alignment
#                     worksheet[f"H{row}"].alignment = wrap_alignment

#                     desc = worksheet[f"C{row}"].value
#                     mitigation = worksheet[f"E{row}"].value
#                     treatment = worksheet[f"H{row}"].value

#                     height_desc = calculate_row_height(desc, column_widths["C"])
#                     height_mit = calculate_row_height(mitigation, column_widths["E"])
#                     height_treat = calculate_row_height(treatment, column_widths["H"])

#                     worksheet.row_dimensions[row].height = max(
#                         height_desc,
#                         height_mit,
#                         height_treat
#                     )

#                 # Header Style
#                 header_alignment = Alignment(
#                     wrap_text=True,
#                     horizontal="center",
#                     vertical="center"
#                 )

#                 header_fill = PatternFill(
#                     start_color="D9E1F2",
#                     end_color="D9E1F2",
#                     fill_type="solid"
#                 )

#                 header_font = Font(bold=True)

#                 for cell in worksheet[1]:
#                     cell.alignment = header_alignment
#                     cell.fill = header_fill
#                     cell.font = header_font

#         output.seek(0)

#         return output

#     except Exception as e:
#         raise e


##-------------------------------------------new function---------------
def get_risk_data_excel(db, dept_id,financial_year):
    
    impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}

    green_value = {"1A", "2A", "3A", "1B", "2B"}
    light_green_value = {"1C","2C", "3B","4A","5A"}
    yellow_value = {"4B", "5B", "3C", "2D","1D"}
    amber_value = {"1E", "2E", "3D", "4C", "5C"}
    red_value = {"3E","4E", "5E", "4D", "5D"}

    green_fill = PatternFill(start_color="008000", end_color="008000", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    amber_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    light_green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

    try:

        query = db.query(RiskRegister).options(
            joinedload(RiskRegister.risk_descriptions)
            .joinedload(RiskDescription.treatments)
            .joinedload(RiskTreatment.action_owner),
            joinedload(RiskRegister.department),
            joinedload(RiskRegister.risk_owner)
        )

        # Department_id filter
        if dept_id:
            query = query.filter(RiskRegister.dept_id == dept_id)
            
        # Financial year filter
        if financial_year:
            query = query.filter(
                RiskRegister.financial_year == financial_year
            )
        else:
            today = datetime.now()

            if today.month >= 4:
                current_fy = f"{today.year}-{today.year + 1}"
            else:
                current_fy = f"{today.year - 1}-{today.year}"

            query = query.filter(
                RiskRegister.financial_year == current_fy)

        query = query.filter(
            RiskRegister.is_deleted == 0,
            RiskRegister.dept_id > 0
        ).order_by(RiskRegister.risk_register_id)

        risks = query.all()

        department_rows = {}
        risk_merge_ranges = {}
        description_merge_ranges = {}
        current_rows = {}

        for risk in risks:

            dept_name = "UNKNOWN"
            if risk.department:
                dept_name = risk.department.dept_short_name

            if dept_name not in department_rows:
                department_rows[dept_name] = []
                risk_merge_ranges[dept_name] = []
                description_merge_ranges[dept_name] = []
                current_rows[dept_name] = 2

            risk_owner_name = risk.risk_owner.log_id if risk.risk_owner else ""

            risk_start_row = current_rows[dept_name]

            descriptions = risk.risk_descriptions if risk.risk_descriptions else [None]

            for desc in descriptions:

                desc_start_row = current_rows[dept_name]

                likelihood = desc.inherent_risk_likelihood_id if desc else None
                impact = desc.inherent_risk_impact_id if desc else None

                current_likelihood = desc.current_risk_likelihood_id if desc else None
                current_impact = desc.current_risk_impact_id if desc else None

                inherent_color_code = ""
                current_color_code = ""

                if likelihood and impact:
                    inherent_color_code = f"{likelihood}{impact_map.get(impact)}"

                if current_likelihood and current_impact:
                    current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"

                treatments = desc.treatments if desc and desc.treatments else [None]

                first_desc = True

                for treatment in treatments:

                    department_rows[dept_name].append({

                        "Risk ID": risk.risk_id if current_rows[dept_name] == risk_start_row else "",
                        "Risk Title": risk.risk_name if current_rows[dept_name] == risk_start_row else "",

                        "Risk Description": desc.risk_description if desc and first_desc else "",
                        "Inherent Risk Level": inherent_color_code if first_desc else "",
                        "Current Mitigation": desc.mitigation if desc and first_desc else "",
                        "Current Risk Level": current_color_code if first_desc else "",

                        "Risk Owner": risk_owner_name if current_rows[dept_name] == risk_start_row else "",
                        "Action Plan": treatment.action_plan if treatment else "",
                        "Action Owner": treatment.action_owner.log_id if treatment and treatment.action_owner else "",
                        "Target Date": treatment.target_date.date() if treatment and treatment.target_date else "",
                        "Action Status": treatment.status.status_name if treatment and treatment.status else ""
                    })

                    first_desc = False
                    current_rows[dept_name] += 1

                desc_end_row = current_rows[dept_name] - 1

                if desc_end_row >= desc_start_row:
                    description_merge_ranges[dept_name].append(
                        (desc_start_row, desc_end_row)
                    )

            risk_end_row = current_rows[dept_name] - 1

            if risk_end_row >= risk_start_row:
                risk_merge_ranges[dept_name].append(
                    (risk_start_row, risk_end_row)
                )

        output = io.BytesIO()

        wrap_alignment = Alignment(wrap_text=True, vertical="center")

        merge_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        with pd.ExcelWriter(output, engine="openpyxl") as writer:

            for dept_name, rows in department_rows.items():

                df = pd.DataFrame(rows)

                sheet_name = dept_name[:31]

                df.to_excel(writer, index=False, sheet_name=sheet_name)

                worksheet = writer.sheets[sheet_name]

                # Color Fill
                for row in range(2, len(df) + 2):

                    cell_d = worksheet[f"D{row}"]
                    cell_f = worksheet[f"F{row}"]

                    if cell_d.value in green_value:
                        cell_d.fill = green_fill
                    elif cell_d.value in amber_value:
                        cell_d.fill = amber_fill
                    elif cell_d.value in yellow_value:
                        cell_d.fill = yellow_fill
                    elif cell_d.value in red_value:
                        cell_d.fill = red_fill
                    elif cell_d.value in light_green_value:
                        cell_d.fill = light_green_fill

                    if cell_f.value in green_value:
                        cell_f.fill = green_fill
                    elif cell_f.value in amber_value:
                        cell_f.fill = amber_fill
                    elif cell_f.value in yellow_value:
                        cell_f.fill = yellow_fill
                    elif cell_f.value in red_value:
                        cell_f.fill = red_fill
                    elif cell_f.value in light_green_value:
                        cell_f.fill = light_green_fill

                # Column Width
                column_widths = {
                    "A": 15,   # Risk ID (A)
                    "B" :20, # Risk Name (B)
                    "C": 50,   # Risk Description (C)
                    "D": 10,   # Inherent Risk Level (D)
                    "E": 70,    # Mitigation (E)
                    "F": 10, # Current risk level (F)
                    "G": 15, # Risk Owner (G)
                    "H": 70, # Treatment (H)
                    "I": 15, # Action Owner (I)
                    "J": 15,  # Target Date (J)
                    "K":15 # Status
                }

                for col, width in column_widths.items():
                    worksheet.column_dimensions[col].width = width

                # Risk Merge
                for start, end in risk_merge_ranges[dept_name]:

                    if start != end:

                        worksheet.merge_cells(f"A{start}:A{end}")
                        worksheet.merge_cells(f"B{start}:B{end}")
                        worksheet.merge_cells(f"G{start}:G{end}")

                    worksheet[f"A{start}"].alignment = merge_alignment
                    worksheet[f"B{start}"].alignment = merge_alignment
                    worksheet[f"G{start}"].alignment = merge_alignment

                # Description Merge
                for start, end in description_merge_ranges[dept_name]:

                    if start != end:

                        worksheet.merge_cells(f"C{start}:C{end}")
                        worksheet.merge_cells(f"D{start}:D{end}")
                        worksheet.merge_cells(f"E{start}:E{end}")
                        worksheet.merge_cells(f"F{start}:F{end}")

                    worksheet[f"C{start}"].alignment = merge_alignment
                    worksheet[f"D{start}"].alignment = merge_alignment
                    worksheet[f"E{start}"].alignment = merge_alignment
                    worksheet[f"F{start}"].alignment = merge_alignment

                # Wrap Text & Dynamic Height
                for row in range(2, len(df) + 2):

                    worksheet[f"C{row}"].alignment = wrap_alignment
                    worksheet[f"E{row}"].alignment = wrap_alignment
                    worksheet[f"H{row}"].alignment = wrap_alignment

                    desc = worksheet[f"C{row}"].value
                    mitigation = worksheet[f"E{row}"].value
                    treatment = worksheet[f"H{row}"].value

                    height_desc = calculate_row_height(desc, column_widths["C"])
                    height_mit = calculate_row_height(mitigation, column_widths["E"])
                    height_treat = calculate_row_height(treatment, column_widths["H"])

                    worksheet.row_dimensions[row].height = max(
                        height_desc,
                        height_mit,
                        height_treat
                    )

                # Header Style
                header_alignment = Alignment(
                    wrap_text=True,
                    horizontal="center",
                    vertical="center"
                )

                header_fill = PatternFill(
                    start_color="D9E1F2",
                    end_color="D9E1F2",
                    fill_type="solid"
                )

                header_font = Font(bold=True)

                for cell in worksheet[1]:
                    cell.alignment = header_alignment
                    cell.fill = header_fill
                    cell.font = header_font

        output.seek(0)

        return output

    except Exception as e:
        raise e
        
    
    
# def get_risk_data_excel_previous(db,dept_id):
#     impact_map = {1:"A",2:"B",3:"C",4:"D",5:"E"}
#     green_value = {"1A", "2A", "3A", "1B", "2B","1C"}
#     yellow_value = {"4A", "5A", "3B", "3C", "2D","2C","1D","1E"}
#     amber_value = {"2E", "3E", "3D", "4C", "4B","5B"}
#     red_value = {"4E", "5E", "4D", "5D", "5C"}
    
#     green_fill = PatternFill(start_color="4CAF50",
#                             end_color="4CAF50",
#                             fill_type="solid")
#     yellow_fill = PatternFill(start_color="FFEB3B",
#                             end_color="FFEB3B",
#                             fill_type="solid")
#     amber_fill = PatternFill(start_color="FF9800",
#                             end_color="FF9800",
#                             fill_type="solid")
#     red_fill = PatternFill(start_color="F44336",
#                             end_color="F44336",
#                             fill_type="solid")
#     try:
#         query = db.query(RiskRegister).options(
#                 joinedload(RiskRegister.risk_descriptions)
#                 .joinedload(RiskDescription.treatments)
#             )
#         if dept_id:
#             query = query.filter(RiskRegister.dept_id == dept_id)
#         query = query.filter(RiskRegister.is_deleted == 0, RiskRegister.dept_id > 0).order_by(RiskRegister.risk_register_id)
#         risks = query.all()

#         department_rows  = {}
#         merge_ranges = {}
#         current_rows = {}

#         #current_row = 2
#         for risk in risks:
#             first_risk = True
#             risk_owner_name = ""
#             dept_name = "UNKNOWN"
#             if risk.risk_owner :
#                 risk_owner_name = risk.risk_owner.log_id

#             if risk.department:
#                 dept_name = risk.department.dept_short_name

#             if dept_name not in department_rows:
#                 department_rows[dept_name] = []
#                 merge_ranges[dept_name] = []
#                 current_rows[dept_name] = 2

#             start_row = current_rows[dept_name]
#             descriptions = risk.risk_descriptions if risk.risk_descriptions else [None]
            
#             for desc in descriptions:

#                 first_desc = True
#                 likelihood = desc.inherent_risk_likelihood_id if desc else None
#                 impact = desc.inherent_risk_impact_id if desc else None
#                 current_likelihood = desc.current_risk_likelihood_id if desc else None
#                 current_impact = desc.current_risk_impact_id if desc else None

#                 inherent_color_str = None
#                 inherent_color_code = None
#                 current_color_str = None
#                 current_color_code = None

#                 if likelihood and impact:
#                     inherent_color_str = get_color(likelihood*impact)
#                     inherent_color_code = f"{likelihood}{impact_map.get(impact)}"
#                 if current_likelihood and current_impact:
#                     current_color_str = get_color(current_likelihood*current_impact)
#                     current_color_code = f"{current_likelihood}{impact_map.get(current_impact)}"

#                 treatments = desc.treatments if desc and desc.treatments else [None]
#                 for treatment in treatments:

#                     department_rows[dept_name].append({
#                         "Risk ID": risk.risk_id if first_risk else "",
#                         "Risk Name" : risk.risk_name if first_risk else "",
#                         "Risk Description": desc.risk_description if desc and first_desc else "",
#                         "Inherent Risk Level" : inherent_color_code if first_desc else "",
#                         "Current Mitigation": desc.mitigation if desc and first_desc else "",
#                         "Current Risk Level" : current_color_code if first_desc else "",
#                         "Risk Owner" : risk_owner_name if first_risk else "",
#                         "Action Owner" : treatment.action_owner.log_id if treatment and treatment.action_owner else "",
#                         "Risk Treatment": treatment.action_plan if treatment else "",
#                         "Target Date" :treatment.target_date.date() if treatment and treatment.target_date else ""
                        
#                     })

#                     first_risk = False
#                     first_desc = False
#                     current_rows[dept_name] += 1
#             end_row = current_rows[dept_name] - 1
#             if end_row >= start_row:
#                 merge_ranges[dept_name].append((start_row, end_row))

#         #df = pd.DataFrame(rows)

#         output = io.BytesIO()
#         wrap_alignment = Alignment(
#             wrap_text=True,
#             vertical="center"
#         )

#         merge_alignment = Alignment(
#             horizontal="center",
#             vertical="center",
#             wrap_text=True
#         )

#         with pd.ExcelWriter(output, engine="openpyxl") as writer:
#             for dept_name, rows in department_rows.items():
#                 df = pd.DataFrame(rows)
#                 sheet_name = dept_name[:31]
#                 df.to_excel(writer, index=False, sheet_name=sheet_name)

#                 workbook = writer.book
#                 worksheet = writer.sheets[sheet_name]
            

#                 # Apply color logic
#                 for row in range(2, len(df) + 2):   # Excel row index starts at 1
#                     cell = worksheet[f"D{row}"]     # Column D = Inherent Risk Level

#                     if cell.value in green_value:
#                         cell.fill = green_fill
#                     elif cell.value in amber_value:
#                         cell.fill = amber_fill
#                     elif cell.value in yellow_value:
#                         cell.fill = yellow_fill
#                     elif cell.value in red_value:
#                         cell.fill = red_fill

#                     cell_f = worksheet[f"F{row}"]     # Column D = Inherent Risk Level

#                     if cell_f.value in green_value:
#                         cell_f.fill = green_fill
#                     elif cell_f.value in amber_value:
#                         cell_f.fill = amber_fill
#                     elif cell_f.value in yellow_value:
#                         cell_f.fill = yellow_fill
#                     elif cell_f.value in red_value:
#                         cell_f.fill = red_fill

#                 # Column Widths
#                 column_widths = {
#                     "A": 15,   # Risk ID (A)
#                     "B" :20, # Risk Name (B)
#                     "C": 50,   # Risk Description (C)
#                     "D": 10,   # Inherent Risk Level (D)
#                     "E": 70,    # Mitigation (E)
#                     "F": 10, # Current risk level (F)
#                     "G": 15, # Risk Owner (G)
#                     "H": 15, # Action Owner (H)
#                     "I": 70, # Treatment (I)
#                     "J": 15  # Target Date (J)
#                 }

#                 for col, width in column_widths.items():
#                     worksheet.column_dimensions[col].width = width

                
#                 for start, end in merge_ranges[dept_name]:
#                     if start != end:

#                         worksheet.merge_cells(f"A{start}:A{end}")  # Risk ID
#                         worksheet.merge_cells(f"B{start}:B{end}")  # Risk Name
#                         # worksheet.merge_cells(f"C{start}:C{end}")  # Risk Description
#                         # worksheet.merge_cells(f"D{start}:D{end}")  # Inherent Risk Level
#                         # worksheet.merge_cells(f"E{start}:E{end}")  # Current Mitigation
#                         # worksheet.merge_cells(f"F{start}:F{end}")  # Current Risk Level
#                         worksheet.merge_cells(f"G{start}:G{end}")  # Risk Owner
                    
#                     worksheet[f"A{start}"].alignment = merge_alignment
#                     worksheet[f"B{start}"].alignment = merge_alignment
#                     # worksheet[f"C{start}"].alignment = merge_alignment
#                     # worksheet[f"D{start}"].alignment = merge_alignment
#                     # worksheet[f"E{start}"].alignment = merge_alignment
#                     # worksheet[f"F{start}"].alignment = merge_alignment
#                     worksheet[f"G{start}"].alignment = merge_alignment

#                 # Wrap Text for all cells
#                 for row in range(2, len(df)+2):
#                     worksheet[f"C{row}"].alignment = wrap_alignment   # Description
#                     worksheet[f"E{row}"].alignment = wrap_alignment   # Mitigation
#                     worksheet[f"I{row}"].alignment = wrap_alignment   # Treatment
                    
#                     desc = worksheet[f"C{row}"].value
#                     mitigation = worksheet[f"E{row}"].value
#                     treatment = worksheet[f"I{row}"].value

#                     height_desc = calculate_row_height(desc, column_widths["C"])
#                     height_mit = calculate_row_height(mitigation, column_widths["E"])
#                     height_treat = calculate_row_height(treatment, column_widths["I"])

#                     worksheet.row_dimensions[row].height = max(height_desc, height_mit, height_treat)

#                 # Wrap Header
#                 header_alignment = Alignment(
#                     wrap_text=True,
#                     horizontal="center",
#                     vertical="center"
#                 )
                
#                 header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
#                 header_font = Font(bold=True)

#                 for cell in worksheet[1]:
#                     cell.alignment = header_alignment
#                     cell.fill = header_fill
#                     cell.font = header_font

#         output.seek(0)

#         return output
#     except Exception as e:
#         raise e



    

#---------
# Get Followups by reference_id
#---------

def get_followups_by_reference_id(db, reference_id):
    try:

        followups = (
            db.query(RiskActionFollowup)
            .options(
                joinedload(RiskActionFollowup.created_user)
                .load_only(User.log_id),
                
                joinedload(RiskActionFollowup.status_master)
                .load_only(Status.status_name)
            )
            .filter(
                RiskActionFollowup.reference_id == reference_id
            )
            .all()
        )

        result = []

        for followup in followups:

            followup_dict = followup.__dict__.copy()

            # remove SQLAlchemy internal state
            followup_dict.pop("_sa_instance_state", None)
            followup_dict.pop("file_data", None)

            # add user name
            followup_dict["risk_owner_name"] = (
                followup.created_user.log_id if followup.created_user else None
            )
            
            # add status name
            followup_dict["risk_status_name"] = (
                followup.status_master.status_name if followup.status_master else None
            )

            # remove relationship object
            followup_dict.pop("created_user", None)
            followup_dict.pop("status_master", None)

            result.append(followup_dict)

        return result

    except Exception as e:
        raise e
    
def get_last_risk_statusbyid(db, risk_id):
        try:
            #Read Risk register based on risk_register_id
            # If all approval done than no need to read history
            risk_register = (
                db.query(RiskRegister)
                .filter(
                    RiskRegister.risk_register_id == risk_id,
                    RiskRegister.risk_function_head_approval_status == 1,
                    RiskRegister.risk_manager_approval_status == 1,
                    RiskRegister.risk_head_approval_status == 1
                )
                .all()
            )
            data = []
            if risk_register:
                return data

            # APPROVAL STATUS NAMES
            # # read status from risk history table last entry
            risk_history = (
                db.query(RiskRegisterHist)
                .filter(
                    RiskRegisterHist.risk_register_id == risk_id,
                    RiskRegisterHist.modified_on.isnot(None),
                    or_(
                        RiskRegisterHist.risk_function_head_approval_status.isnot(None),
                        RiskRegisterHist.risk_manager_approval_status.isnot(None),
                        RiskRegisterHist.risk_head_approval_status.isnot(None)
                    ),
                    or_(
                        RiskRegisterHist.risk_function_head_approval_status == -1,
                        RiskRegisterHist.risk_manager_approval_status == -1,
                        RiskRegisterHist.risk_head_approval_status == -1
                    )
                )
                .order_by(RiskRegisterHist.modified_on.desc())
                .first()
            )
            function_head_approval_status_name = ""
            if risk_history != None:
                if risk_history.risk_function_head_approval_status != None:
                    function_head_approval_status_name = (
                        get_approval_status_name(
                            risk_history.risk_function_head_approval_status
                        )
                    )

            risk_manager_approval_status_name = ""
            if risk_history != None:
                if risk_history.risk_manager_approval_status != None:
                    risk_manager_approval_status_name = (
                        get_approval_status_name(
                            risk_history.risk_manager_approval_status
                        )
                    )
            
            risk_head_approval_status_name = ""
            if risk_history != None:
                if risk_history.risk_head_approval_status != None:
                    risk_head_approval_status_name = (
                        get_approval_status_name(
                            risk_history.risk_head_approval_status
                        )
                    )
            result = []
            if function_head_approval_status_name != "":
                risk_last_status = {
                    **to_dict(risk_history),
                    "function_head_approval_status_name":function_head_approval_status_name,
                    "risk_manager_approval_status_name" : risk_manager_approval_status_name,
                    "risk_head_approval_status_name": risk_head_approval_status_name
                }
                result.append(risk_last_status)
            return result
        except Exception as e:
            raise e
        

        
# Using FY copy data 

def copy_risks_fy(
    db,
    source_fy: str,
    destination_fy: str,
    current_user_id: int,
    department_short_name: str = None,
    risk_id: str = None
):
    try:

        # Preload Department Counters
        departments = (
            db.query(Department)
            .all()
        )

        dept_counters = {}

        for dept in departments:
            dept_counters[dept.id] = {
                "last_number": dept.last_risk_number,
                "short_name": dept.dept_short_name,
                "obj": dept
            }

        # Cache Columns
        risk_columns = [
            c.name
            for c in RiskRegister.__table__.columns
            if c.name not in (
                "risk_register_id",
                "created_on",
                "modified_on",
                "modified_by"
            )
        ]

        desc_columns = [
            c.name
            for c in RiskDescription.__table__.columns
            if c.name not in (
                "risk_description_id",
                "created_on",
                "modified_on",
                "modified_by"
            )
        ]

        treatment_columns = [
            c.name
            for c in RiskTreatment.__table__.columns
            if c.name not in (
                "risk_treatment_id",
                "created_on",
                "modified_on",
                "modified_by"
            )
        ]

        total_copied = 0
        total_descriptions = 0
        total_treatments = 0

        offset = 0
        batch_size = 100
        
        
        # Validate Source FY
        source_risk_exists = (
            db.query(RiskRegister)
            .filter(
                RiskRegister.financial_year == source_fy,
                RiskRegister.is_deleted == 0
            )
            .first()
        )

        if not source_risk_exists:
            raise Exception(f"No risks found for source financial year '{source_fy}'")
        
         # Validate Destination year data already exists
        destination_risk_exists = (
            db.query(RiskRegister)
            .filter(
                RiskRegister.financial_year == destination_fy,
                RiskRegister.is_deleted == 0
            )
            .first()
        )

        if destination_risk_exists:
            raise Exception(f"Data already exists '{destination_fy}'")
        
        
        # Validate Department        
        if department_short_name:

            department_exists = (
                db.query(Department).filter(Department.dept_short_name == department_short_name).first())

            if not department_exists:
                raise Exception(f"Department '{department_short_name}' not found")

        # Validate Risk ID
        if risk_id:
            risk_exists = (
                db.query(RiskRegister)
                .filter(
                    RiskRegister.risk_id == risk_id,
                    RiskRegister.financial_year == source_fy,
                    RiskRegister.is_deleted == 0
                )
                .first()
            )

            if not risk_exists:
                raise Exception(f"Risk '{risk_id}' not found in FY {source_fy}")

        while True:

            query = (
                db.query(RiskRegister)
                .options(
                    joinedload(RiskRegister.risk_descriptions)
                    .joinedload(RiskDescription.treatments)
                )
                .filter(
                    RiskRegister.financial_year == source_fy,
                    RiskRegister.is_deleted == 0
                )
            )

            # Department
            if department_short_name:

                query = (query.join(Department,RiskRegister.dept_id == Department.id)
                    .filter(Department.dept_short_name == department_short_name))

            # Risk ID 
            if risk_id:
                query = (query.filter(RiskRegister.risk_id == risk_id))

            risks = (
                query
                .order_by(RiskRegister.risk_register_id)
                .offset(offset)
                .limit(batch_size)
                .all()
            )

            if not risks:
                break

            # ------------------------------------------
            # Process Batch
            # ------------------------------------------
            for old_risk in risks:

                dept_info = dept_counters.get(
                    old_risk.dept_id
                )

                if not dept_info:
                    raise Exception(
                        f"Department not found: {old_risk.dept_id}"
                    )

                dept_info["last_number"] += 1

                new_risk_code = (
                    f"{dept_info['short_name']}-"
                    f"{str(dept_info['last_number']).zfill(4)}"
                )

                # Risk Register
                new_risk = RiskRegister()

                for col in risk_columns:
                    setattr(
                        new_risk,
                        col,
                        getattr(old_risk, col)
                    )

                new_risk.risk_id = new_risk_code
                new_risk.financial_year = destination_fy

                # reset columns
                new_risk.risk_function_head_approval_status = None
                new_risk.risk_function_head_approval_remark = None
                new_risk.risk_function_head_approval_on = None
                new_risk.risk_function_head_approval_by = None

                new_risk.risk_head_approval_status = None
                new_risk.risk_head_approval_remark = None
                new_risk.risk_head_approved_on = None
                new_risk.risk_head_approval_by = None

                new_risk.risk_manager_approval_status = None
                new_risk.risk_manager_approval_remark = None
                new_risk.risk_manager_approved_on = None
                new_risk.risk_manager_approval_by = None
                
                new_risk.risk_status = None
                new_risk.risk_progress = None

                new_risk.created_by = current_user_id
                new_risk.created_on = datetime.now()

                db.add(new_risk)
                db.flush()

                total_copied += 1

                # Descriptions
                for old_desc in old_risk.risk_descriptions:

                    total_descriptions += 1

                    new_desc = RiskDescription()

                    for col in desc_columns:
                        setattr(
                            new_desc,
                            col,
                            getattr(old_desc, col)
                        )

                    new_desc.risk_register_id = (
                        new_risk.risk_register_id
                    )

                    new_desc.risk_id = (
                        new_risk.risk_id
                    )

                    new_desc.created_by = current_user_id
                    new_desc.created_on = datetime.now()

                    db.add(new_desc)
                    db.flush()

                    # Treatments
                    for old_treatment in old_desc.treatments:

                        total_treatments += 1

                        new_treatment = RiskTreatment()

                        for col in treatment_columns:
                            setattr(
                                new_treatment,
                                col,
                                getattr(old_treatment, col)
                            )

                        new_treatment.risk_register_id = (
                            new_risk.risk_register_id
                        )

                        new_treatment.risk_description_id = (
                            new_desc.risk_description_id
                        )

                        new_treatment.risk_id = (
                            new_risk.risk_id
                        )

                        new_treatment.target_date = None
                        new_treatment.action_status_id = None
                        new_treatment.progress = None
                        
                        new_treatment.next_followup_date = None


                        new_treatment.created_by = current_user_id
                        new_treatment.created_on = datetime.now()

                        db.add(new_treatment)

            db.commit()

            print(
                f"Processed batch "
                f"{offset} - {offset + len(risks)}"
            )

            offset += batch_size

        # Update Department Counters Once
        for dept_id, dept_info in dept_counters.items():

            dept_info["obj"].last_risk_number = (
                dept_info["last_number"]
            )

        db.commit()

        return {
            "status": True,
            "message": "Risk copy completed",
            "source_financial_year": source_fy,
            "destination_financial_year": destination_fy,
            "department_short_name": department_short_name,
            "risk_id": risk_id,
            "total_risks_copied": total_copied,
            "total_descriptions_copied": total_descriptions,
            "total_treatments_copied": total_treatments
        }

    except Exception as e:
        db.rollback()
        raise e
    
    
## Get distinct Dept for Risk Owner and Risk CO Owner

def get_my_departments(db: Session, current_user):
    departments = (
        db.query(
            Department.id,
            Department.dept_name
        )
        .join(
            RiskRegister,
            RiskRegister.dept_id == Department.id
        )
        .filter(
            RiskRegister.is_deleted == 0,
            or_(
                RiskRegister.risk_owner_id == current_user["id"],
                RiskRegister.risk_co_owner_id == current_user["id"]
            )
        )
        .distinct()
        .order_by(Department.dept_name)
        .all()
    )

    return [
        {
            "dept_id": dept.id,
            "dept_name": dept.dept_name
        }
        for dept in departments
    ]